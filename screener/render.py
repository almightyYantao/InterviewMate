"""渲染：完整报告 / 面试记录表（不含答案）/ docx。"""
from __future__ import annotations

from collections import OrderedDict

from .parsing import unmask
from .schemas import DECISION_LABEL, SECTION_ORDER, Report

ICON = {"match": "✅", "partial": "⚠️", "missing": "❌", "unknown": "❓"}
SEV = {"high": "🔴 高", "medium": "🟠 中", "low": "⚪ 低"}
LEVEL = {"must": "必须", "plus": "加分", "duty": "职责"}


def _cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", " ")


def _sections(r: Report) -> "OrderedDict[str, list]":
    secs: OrderedDict[str, list] = OrderedDict((s, []) for s in SECTION_ORDER)
    for q in r.questions:
        secs[q.section].append(q)
    return OrderedDict((k, v) for k, v in secs.items() if v)


def _source(q, r: Report) -> str:
    if q.source.type == "resume":
        return f"简历原文（第 {', '.join(map(str, q.source.lines))} 行）：「{q.source.quote}」"
    if q.source.type == "jd":
        reqs = {x.id: x.text for x in r.jd.requirements}
        ids = q.linked_requirements
        return "JD 条目：" + ("；".join(f"{i} {reqs.get(i, '')}" for i in ids) if ids else q.source.quote)
    titles = {d.id: d.title for d in r.doubts}
    ids = q.linked_doubts
    return "疑点：" + ("；".join(f"{i} {titles.get(i, '')}" for i in ids) if ids else q.source.quote)


def markdown(r: Report, pii: dict[str, str] | None = None, with_answers: bool = True) -> str:
    c = r.conclusion
    reqs = {x.id: x for x in r.jd.requirements}
    total = sum(q.minutes for q in r.questions)
    out: list[str] = []
    title = "简历筛选与一面准备" if with_answers else "一面面试记录表"
    out.append(f"# {r.candidate} · {r.position}{title}\n")
    out.append("> ⚠️ AI 生成，需人工复核。系统只给出建议，是否推进由面试官决定。"
               "出生日期仅用于时间线自洽校验，年龄、性别、婚育、籍贯不参与评分和出题。\n")

    if with_answers:
        out.append("## 一、筛选结论\n")
        out.append(f"**结论：{DECISION_LABEL[c.decision]}** ｜ 匹配分 {c.score}/10\n")
        out.append(c.summary + "\n")
        if c.veto:
            out.append("**一票否决项**（结论上限为「有条件推进」）：" +
                       "；".join(f"{v} {reqs[v].text if v in reqs else ''}" for v in c.veto) + "\n")
        out.append("**必须核实**：" + "、".join(c.must_verify) + "\n")

        out.append("## 二、简历疑点\n")
        out.append("| # | 疑点 | 级别 | 原文证据 | 建议核实问题 |\n|---|---|---|---|---|")
        order = {"high": 0, "medium": 1, "low": 2}
        for d in sorted(r.doubts, key=lambda d: order[d.severity]):
            ev = "<br>".join(f"L{e.line}「{_cell(e.quote)}」" for e in d.evidence) or "—"
            tag = "规则" if d.source == "rule" else "语义"
            out.append(f"| {d.id} | {_cell(d.title)}（{tag}） | {SEV[d.severity]} | {ev} | {_cell(d.verify_question)} |")
        out.append("")

        out.append("## 三、JD 匹配度\n")
        out.append("| JD 要求 | 级别 / 权重 | 匹配 | 证据 / 备注 |\n|---|---|---|---|")
        for m in r.match:
            req = reqs.get(m.req_id)
            text = f"{m.req_id} {req.text}" if req else m.req_id
            if req and req.emphasized:
                text = f"**{text}**"
            ev = "<br>".join(f"L{e.line}「{_cell(e.quote)}」" for e in m.evidence)
            note = _cell(m.reason) + (f"<br>{ev}" if ev else "")
            lvl = f"{LEVEL[req.level]} / {req.weight}" if req else ""
            out.append(f"| {_cell(text)} | {lvl} | {ICON[m.verdict]} | {note} |")
        out.append("")

    out.append(f"## {'四、' if with_answers else ''}面试题{'与参考答案' if with_answers else ''}（共 {total} 分钟）\n")
    for sec, qs in _sections(r).items():
        out.append(f"### {sec}（{sum(q.minutes for q in qs)} 分钟）\n")
        for q in qs:
            out.append(f"**{q.id}. {q.question}**（{q.minutes} 分钟）\n")
            out.append(f"> 来源：{_source(q, r)}\n")
            if with_answers:
                key = set(q.key_points)
                pts = "；".join(f"**{p}**" if p in key else p for p in q.answer_points)
                out.append(f"- 参考答案要点：{pts}")
                out.append(f"- 追问：{'；'.join(q.follow_ups)}")
                out.append(f"- 🚩 红旗：{'；'.join(q.red_flags)}")
                out.append(f"- 评分：优秀 —— {q.rubric.excellent}；合格 —— {q.rubric.pass_}；不合格 —— {q.rubric.fail}\n")
            else:
                if q.follow_ups:
                    out.append(f"- 追问：{'；'.join(q.follow_ups)}")
                out.append("- 评分：☐ 优秀　☐ 合格　☐ 不合格")
                out.append("- 记录：\n\n&nbsp;\n")

    if with_answers:
        out.append("## 五、评价建议\n")
        out.append("- 通过信号：" + "；".join(c.pass_signals))
        out.append("- 不通过信号：" + "；".join(c.fail_signals))
        out.append("")
        if r.validation:
            out.append("<details><summary>校验层日志</summary>\n")
            out.extend(f"- {_cell(x)}" for x in r.validation)
            out.append("\n</details>\n")
    else:
        out.append("## 面试结论\n\n☐ 通过　☐ 待定　☐ 不通过\n\n评价：\n")

    text = "\n".join(out)
    return unmask(text, pii) if pii else text


def docx_bytes(md: str) -> bytes:
    """把上面的 Markdown 粗略转成 Word（标题 / 段落 / 列表 / 表格）。"""
    import io
    import re

    import docx

    d = docx.Document()
    lines = md.splitlines()
    i = 0
    strip = lambda s: re.sub(r"\*\*|<br>|<details>|</details>|<summary>|</summary>|&nbsp;", lambda m: "\n" if m.group(0) == "<br>" else "", s)
    while i < len(lines):
        l = lines[i]
        if l.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= {"-"}:
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not set(lines[i].replace("|", "").strip()) <= {"-"}:
                    rows.append([strip(c.strip()).replace("\\|", "|") for c in re.split(r"(?<!\\)\|", lines[i])[1:-1]])
                i += 1
            t = d.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r_, row in enumerate(rows):
                for c_, val in enumerate(row[: len(rows[0])]):
                    t.cell(r_, c_).text = val
            continue
        m = re.match(r"^(#{1,3}) (.*)", l)
        if m:
            d.add_heading(strip(m.group(2)), level=len(m.group(1)))
        elif l.startswith("- "):
            d.add_paragraph(strip(l[2:]), style="List Bullet")
        elif l.startswith("> "):
            d.add_paragraph(strip(l[2:]), style="Intense Quote")
        elif strip(l).strip():
            d.add_paragraph(strip(l))
        i += 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
