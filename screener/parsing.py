"""Stage 0 · 文档解析与清洗。"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

SCAN_CHARS_PER_PAGE = 50


@dataclass
class Doc:
    clean_text: str
    lines: list[str]                       # lines[i] 对应行号 i+1
    pages: list[int] = field(default_factory=list)   # 每行所在页码
    bold_lines: set[int] = field(default_factory=set)  # docx 中加粗的行号
    scanned: bool = False
    pdf_bytes: bytes | None = None         # 扫描版直接交给 Claude 读
    format_mess: float = 0.0               # 编号格式混乱度 0~1
    pii: dict[str, str] = field(default_factory=dict)  # 占位符 -> 原文

    def numbered(self) -> str:
        return "\n".join(f"{i + 1:>4}| {l}" for i, l in enumerate(self.lines))


# ---------- 读取 ----------


def _is_watermark_char(obj: dict) -> bool:
    """招聘平台水印一般是浅灰色字符，正文是黑色或彩色。"""
    if obj.get("object_type") != "char":
        return False
    c = obj.get("non_stroking_color")
    if isinstance(c, (tuple, list)) and len(c) == 3:
        r, g, b = c
        return abs(r - g) < 0.02 and abs(g - b) < 0.02 and 0.2 < r < 0.95
    if isinstance(c, (tuple, list)) and len(c) == 1:
        return 0.2 < c[0] < 0.95
    return False


def _is_overlay_char(obj: dict) -> bool:
    """招聘平台叠加的水印层：没嵌入的标准字体（如 Helvetica，字体名没有「XXXXXX+」子集前缀）或斜排。"""
    if obj.get("object_type") != "char":
        return False
    return not obj.get("upright", True) or "+" not in (obj.get("fontname") or "")


def clean_page(page):
    """去掉招聘平台水印后的页面。先按浅灰色去；正文本身是深灰色、被删掉大半时，改按叠加字体去。"""
    by_color = page.filter(lambda o: not _is_watermark_char(o))
    raw = sum(1 for o in page.chars if not o["text"].isspace())
    kept = sum(1 for o in page.chars if not o["text"].isspace() and not _is_watermark_char(o))
    if raw > 300 and kept < 0.25 * raw:
        return page.filter(lambda o: not _is_overlay_char(o))
    return by_color


def _read_pdf(path: Path) -> tuple[list[str], list[int], bool]:
    import pdfplumber

    lines, pages, chars_per_page = [], [], []
    with pdfplumber.open(path) as pdf:
        for no, page in enumerate(pdf.pages, 1):
            text = clean_page(page).extract_text(x_tolerance=1.5, y_tolerance=7) or ""
            chars_per_page.append(len(re.sub(r"\s", "", text)))
            for l in text.splitlines():
                lines.append(l)
                pages.append(no)
    scanned = bool(chars_per_page) and max(chars_per_page) < SCAN_CHARS_PER_PAGE
    return lines, pages, scanned


def _read_docx(path: Path) -> tuple[list[str], set[int]]:
    import docx

    d = docx.Document(str(path))
    lines, bold = [], set()
    for p in d.paragraphs:
        if not p.text.strip():
            continue
        lines.append(p.text)
        runs = [r for r in p.runs if r.text.strip()]
        if runs and all(r.bold for r in runs):
            bold.add(len(lines))
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return lines, bold


def _read_markdown(text: str) -> tuple[list[str], set[int]]:
    lines, bold = [], set()
    for l in text.splitlines():
        if not l.strip():
            continue
        if re.search(r"\*\*[^*]+\*\*", l):
            bold.add(len(lines) + 1)
            l = l.replace("**", "")
        lines.append(l)
    return lines, bold


# ---------- 清洗 ----------

_TOKEN = re.compile(r"\S{21,}")
_EMAIL_LIKE = re.compile(r"@|https?://|www\.")


def remove_watermark_tokens(lines: list[str]) -> list[str]:
    """删除长度 > 20、无空格、出现 ≥3 次的 token（文本层面的水印）。"""
    toks = {t for l in lines for t in _TOKEN.findall(l) if not _EMAIL_LIKE.search(t)}
    text = "\n".join(lines)
    bad = {t for t in toks if text.count(t) >= 3}
    # 只保留最短的那个（前后粘了别的字符的变体不重复处理）
    bad = {t for t in bad if not any(o != t and o in t for o in bad)}
    if not bad:
        return lines
    out = []
    for l in lines:
        for t in bad:
            l = l.replace(t, "")
        out.append(l)
    return out


def dedupe_headers(lines: list[str], pages: list[int]) -> tuple[list[str], list[int]]:
    """各页首尾行如果在多页重复出现，视为页眉页脚删除。"""
    if not pages or max(pages) < 3:
        return lines, pages
    edges: Counter[str] = Counter()
    for p in set(pages):
        idx = [i for i, pg in enumerate(pages) if pg == p]
        for i in idx[:1] + idx[-1:]:
            edges[lines[i].strip()] += 1
    repeated = {s for s, n in edges.items() if s and n >= max(3, len(set(pages)) // 2)}

    def is_header(s: str) -> bool:
        return any(SequenceMatcher(None, s, r).ratio() > 0.9 for r in repeated)

    keep = [i for i, l in enumerate(lines) if not (l.strip() and is_header(l.strip()))]
    return [lines[i] for i in keep], [pages[i] for i in keep]


_NUM_STYLES = {
    "cn": re.compile(r"^\s*[一二三四五六七八九十]+[、．.]"),
    "paren": re.compile(r"^\s*\d+[)）]"),
    "dot": re.compile(r"^\s*\d+[．.、](?!\d)"),
    "alpha": re.compile(r"^\s*[A-H][．.、]"),
    "bullet": re.compile(r"^\s*[•●○■□◆▪•]"),
}


def format_mess(lines: list[str]) -> float:
    """编号格式混乱度：使用的编号风格种数 / 5。"""
    used = {k for l in lines for k, p in _NUM_STYLES.items() if p.match(l)}
    return round(len(used) / len(_NUM_STYLES), 2)


# ---------- 脱敏 ----------

_PHONE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_IDCARD = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")


def mask_pii(lines: list[str]) -> tuple[list[str], dict[str, str]]:
    mapping: dict[str, str] = {}
    rev: dict[str, str] = {}

    def sub(kind: str):
        def f(m: re.Match) -> str:
            v = m.group(0)
            if v not in rev:
                ph = f"[{kind}_{sum(k.startswith(f'[{kind}_') for k in mapping) + 1}]"
                rev[v] = ph
                mapping[ph] = v
            return rev[v]
        return f

    out = []
    for l in lines:
        l = _IDCARD.sub(sub("身份证"), l)
        l = _EMAIL.sub(sub("邮箱"), l)
        l = _PHONE.sub(sub("手机"), l)
        out.append(l)
    return out, mapping


def unmask(text: str, mapping: dict[str, str]) -> str:
    for ph, v in mapping.items():
        text = text.replace(ph, v)
    return text


# ---------- 入口 ----------


def load(path: str | Path, mask: bool = True) -> Doc:
    path = Path(path)
    ext = path.suffix.lower()
    pages: list[int] = []
    bold: set[int] = set()
    scanned = False
    pdf_bytes = None
    if ext == ".pdf":
        lines, pages, scanned = _read_pdf(path)
        if scanned:
            pdf_bytes = path.read_bytes()
    elif ext == ".docx":
        lines, bold = _read_docx(path)
    elif ext in {".md", ".txt", ""}:
        lines, bold = _read_markdown(path.read_text(encoding="utf-8"))
    else:
        raise ValueError(f"暂不支持的格式：{ext}（支持 pdf / docx / md / txt；图片请先转 PDF）")
    return from_lines(lines, pages, bold, scanned, pdf_bytes, mask)


def load_text(text: str, mask: bool = False) -> Doc:
    lines, bold = _read_markdown(text)
    return from_lines(lines, [], bold, False, None, mask)


def from_lines(lines, pages, bold, scanned, pdf_bytes, mask) -> Doc:
    lines = remove_watermark_tokens(lines)
    if pages:
        lines, pages = dedupe_headers(lines, pages)
    lines = [re.sub(r"[ \t　]+", " ", l).strip() for l in lines]
    keep = [i for i, l in enumerate(lines) if l]
    lines = [lines[i] for i in keep]
    pages = [pages[i] for i in keep] if pages else []
    mapping: dict[str, str] = {}
    if mask:
        lines, mapping = mask_pii(lines)
    return Doc(
        clean_text="\n".join(lines),
        lines=lines,
        pages=pages,
        bold_lines=bold,
        scanned=scanned,
        pdf_bytes=pdf_bytes,
        format_mess=format_mess(lines),
        pii=mapping,
    )
